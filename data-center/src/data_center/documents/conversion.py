import fcntl
import hashlib
import json
import os
from importlib.metadata import version
from pathlib import Path

from data_center_protocol import DocumentInput
from docling_core.transforms.chunker.hierarchical_chunker import (
    ChunkingDocSerializer,
    ChunkingSerializerProvider,
)
from docling_core.transforms.serializer.markdown import MarkdownTableSerializer
from pydantic import BaseModel, ConfigDict, Field

SUPPORTED = {".pdf", ".docx", ".xlsx", ".pptx", ".md", ".markdown", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}


class _MarkdownSerializerProvider(ChunkingSerializerProvider):
    def get_serializer(self, doc):
        return ChunkingDocSerializer(doc=doc, table_serializer=MarkdownTableSerializer())


def identity_hash(value: dict) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                    separators=(",", ":")).encode()).hexdigest()


def _artifacts_hash(root: Path | None) -> str | None:
    if root is None:
        return None
    files = {}
    for path in sorted(root.rglob("*")):
        if path.is_file() and ".cache" not in path.parts:
            digest = hashlib.sha256()
            with path.open("rb") as stream:
                for data in iter(lambda: stream.read(1024 * 1024), b""):
                    digest.update(data)
            files[str(path.relative_to(root))] = digest.hexdigest()
    if not files:
        raise ValueError("Docling模型目录没有文件")
    return identity_hash({"artifacts": files})


class ConversionPolicy(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    max_file_bytes: int = Field(default=10 * 1024 * 1024, gt=0)
    chunk_tokens: int = Field(default=512, ge=128, le=4096)
    artifacts_path: Path | None = None
    tokenizer_cache: Path | None = None
    tesseract_command: str = "tesseract"
    force_full_page_ocr: bool = False


class ConvertedChunk(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    chunk_id: str
    text: str
    text_sha256: str
    start_line: int
    end_line: int
    locator: dict


class ConvertedDocument(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    input: DocumentInput
    revision_id: str
    source_sha256: str
    content_sha256: str
    config_sha256: str
    converter_version: str
    original_path: Path
    converted_path: Path
    structured_path: Path
    structured_sha256: str
    text: str
    chunks: list[ConvertedChunk]


def _atomic_write(path: Path, raw: bytes) -> None:
    pending = path.with_suffix(path.suffix + ".pending")
    with pending.open("wb") as stream:
        stream.write(raw)
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(pending, path)


def _supplement_native_metadata(document, original: Path) -> None:
    from docling_core.types.doc import DocItemLabel

    if original.suffix.casefold() == ".pptx":
        from pptx import Presentation

        presentation = Presentation(original)
        for slide_number, slide in enumerate(presentation.slides, 1):
            for shape in slide.shapes:
                properties = shape._element.xpath(".//p:cNvPr")
                for item in properties:
                    description = item.get("descr")
                    if description:
                        text = f"Slide {slide_number}, {shape.name}: {description}"
                        document.add_text(label=DocItemLabel.TEXT, text=text)


def _locator(document, chunk) -> dict:
    items = []
    worksheets = set()
    for item in chunk.meta.doc_items:
        items.append({"ref": item.self_ref, "label": str(item.label.value),
                      "provenance": [entry.model_dump(mode="json") for entry in item.prov]})
        parent = item.parent
        while parent is not None:
            node = parent.resolve(document)
            if getattr(node, "label", None) is not None and node.label.value == "sheet":
                worksheets.add(node.name)
            parent = node.parent
    result = {"granularity": "docling_items", "doc_items": items}
    if worksheets:
        result["worksheets"] = sorted(worksheets)
    return result


def verify_converted(document: ConvertedDocument) -> None:
    for path, digest in ((document.original_path, document.source_sha256),
                         (document.converted_path, document.content_sha256),
                         (document.structured_path, document.structured_sha256)):
        if path.is_symlink() or hashlib.sha256(path.read_bytes()).hexdigest() != digest:
            raise ValueError("文档文件 SHA-256 不一致")
    raw = document.converted_path.read_text(encoding="utf-8")
    if raw != document.text:
        raise ValueError("派生文档内容不一致")
    from docling_core.types.doc import DoclingDocument, RefItem

    structured = DoclingDocument.model_validate_json(document.structured_path.read_text())
    lines = raw.splitlines(keepends=True)
    covered = set()
    previous_end = 0
    for chunk in document.chunks:
        if chunk.start_line <= previous_end or chunk.end_line > len(lines):
            raise ValueError("文档片段位置越界或重复")
        content = "".join(lines[chunk.start_line - 1:chunk.end_line])
        if content != chunk.text or hashlib.sha256(content.encode()).hexdigest() != chunk.text_sha256:
            raise ValueError("文档片段内容与全文位置不一致")
        expected = identity_hash({"revision_id": document.revision_id, "config_sha256": document.config_sha256,
                                  "start_line": chunk.start_line, "end_line": chunk.end_line,
                                  "text_sha256": chunk.text_sha256, "locator": chunk.locator})
        if chunk.chunk_id != expected:
            raise ValueError("文档片段身份不一致")
        for item in chunk.locator["doc_items"]:
            original = RefItem(cref=item["ref"]).resolve(structured)
            if [entry.model_dump(mode="json") for entry in original.prov] != item["provenance"]:
                raise ValueError("文档原件位置不一致")
        covered.update(range(chunk.start_line - 1, chunk.end_line))
        previous_end = chunk.end_line
    if {index for index, line in enumerate(lines) if line.strip()} - covered:
        raise ValueError("文档片段未覆盖全文")


def convert_document(input: DocumentInput, raw: bytes, output_root: Path,
                     policy: ConversionPolicy) -> ConvertedDocument:
    if not raw or len(raw) > policy.max_file_bytes:
        raise ValueError("原件字节数无效")
    source_sha256 = hashlib.sha256(raw).hexdigest()
    if input.resource.sha256 is None or source_sha256 != input.resource.sha256:
        raise ValueError("原件 SHA-256 不一致")
    suffix = Path(input.resource.resource_id).suffix.casefold() or Path(input.title).suffix.casefold()
    if suffix not in SUPPORTED:
        raise ValueError("文档格式未获支持")
    converter_version = f"docling-{version('docling')}/docling-core-{version('docling-core')}/tiktoken-{version('tiktoken')}"
    revision_id = identity_hash({"source_id": input.source_id, "record_id": input.record_id,
                                "source_revision": input.source_revision, "source_sha256": source_sha256})
    config_sha256 = identity_hash({"policy": policy.model_dump(mode="json"),
                                  "converter_version": converter_version, "language": input.language,
                                  "model_artifacts_sha256": _artifacts_hash(policy.artifacts_path),
                                  "image_ocr_mode": "full_page",
                                  "tokenizer": "cl100k_base", "native_metadata": "pptx-description-v1",
                                  "table_serializer": "docling-markdown-v1",
                                  "worksheet_chunks": "isolated-v1"})
    directory = output_root / revision_id / config_sha256
    directory.mkdir(parents=True, exist_ok=True)
    metadata = directory / "conversion.json"
    with (directory / "conversion.lock").open("a+b") as lock:
        fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
        if metadata.exists():
            cached = ConvertedDocument.model_validate_json(metadata.read_text())
            if cached.input != input or cached.config_sha256 != config_sha256:
                raise ValueError("已有转换身份不一致")
            verify_converted(cached)
            return cached
        original = directory / ("original" + suffix)
        _atomic_write(original, raw)
        import tiktoken
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import (
            OcrMode,
            PdfPipelineOptions,
            TesseractCliOcrOptions,
        )
        from docling.document_converter import DocumentConverter, ImageFormatOption, PdfFormatOption
        from docling_core.transforms.chunker.hybrid_chunker import HybridChunker
        from docling_core.transforms.chunker.tokenizer.openai import OpenAITokenizer

        cache = policy.tokenizer_cache
        if cache is None:
            configured = os.environ.get("TIKTOKEN_CACHE_DIR")
            if not configured:
                raise ValueError("缺少项目 tokenizer 缓存目录")
            cache = Path(configured)
        if os.environ.get("TIKTOKEN_CACHE_DIR") != str(cache):
            raise ValueError("tokenizer 缓存目录与宿主配置不一致")
        options = {}
        if suffix in {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
            if policy.artifacts_path is None or not policy.artifacts_path.is_dir():
                raise ValueError("缺少已下载的 Docling 模型目录")
            pipeline = PdfPipelineOptions(artifacts_path=policy.artifacts_path,
                                          ocr_options=TesseractCliOcrOptions(
                                              lang=[input.language], tesseract_cmd=policy.tesseract_command,
                                              mode=(OcrMode.FULL_PAGE if policy.force_full_page_ocr or suffix != ".pdf"
                                                    else OcrMode.DEFAULT)))
            options = {InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline),
                       InputFormat.IMAGE: ImageFormatOption(pipeline_options=pipeline)}
        result = DocumentConverter(format_options=options).convert(original)
        if result.status.value != "success":
            raise ValueError(f"Docling 转换未完成：{result.status.value}")
        document = result.document
        _supplement_native_metadata(document, original)
        tokenizer = OpenAITokenizer(tokenizer=tiktoken.get_encoding("cl100k_base"),
                                    max_tokens=policy.chunk_tokens)
        chunker = HybridChunker(tokenizer=tokenizer, merge_peers=suffix != ".xlsx",
                                serializer_provider=_MarkdownSerializerProvider())
        chunks, parts, next_line = [], [], 1
        for chunk in chunker.chunk(dl_doc=document):
            locator = _locator(document, chunk)
            text = chunker.contextualize(chunk=chunk).strip()
            if locator.get("worksheets"):
                text = "Worksheet: " + ", ".join(locator["worksheets"]) + "\n" + text
            if not text:
                continue
            text += "\n"
            end_line = next_line + len(text.splitlines()) - 1
            digest = hashlib.sha256(text.encode()).hexdigest()
            chunk_id = identity_hash({"revision_id": revision_id, "config_sha256": config_sha256,
                                      "start_line": next_line, "end_line": end_line,
                                      "text_sha256": digest, "locator": locator})
            chunks.append(ConvertedChunk(chunk_id=chunk_id, text=text, text_sha256=digest,
                                         start_line=next_line, end_line=end_line, locator=locator))
            parts.append(text)
            next_line = end_line + 1
        if not chunks:
            raise ValueError("文档没有有效文字片段")
        text = "".join(parts)
        structured_raw = document.model_dump_json(indent=2).encode()
        converted_path, structured_path = directory / "document.md", directory / "document.docling.json"
        _atomic_write(converted_path, text.encode())
        _atomic_write(structured_path, structured_raw)
        converted = ConvertedDocument(input=input, revision_id=revision_id, source_sha256=source_sha256,
                                       content_sha256=hashlib.sha256(text.encode()).hexdigest(),
                                       config_sha256=config_sha256, converter_version=converter_version,
                                       original_path=original, converted_path=converted_path,
                                       structured_path=structured_path,
                                       structured_sha256=hashlib.sha256(structured_raw).hexdigest(),
                                       text=text, chunks=chunks)
        verify_converted(converted)
        _atomic_write(metadata, converted.model_dump_json(indent=2).encode())
        return converted
