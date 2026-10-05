import asyncio
import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from data_center_protocol import (
    AuthorizationPort,
    AuthorizedScope,
    DataEvidenceRef,
    DataQuery,
    DataSnapshot,
    DocumentChunk,
    DocumentIngest,
    DocumentSearchHit,
    DocumentSearchResult,
    ModelGateway,
    ModelInvocation,
    SourceReadPort,
)
from integration_protocol import CallContext
from openpyxl import load_workbook
from pptx import Presentation
from pydantic_core import to_jsonable_python
from sqlalchemy import Engine, select, tuple_
from sqlalchemy.orm import Session

from data_center.database import authorized_connection
from data_center.documents.conversion import (
    ConversionPolicy,
    ConvertedDocument,
    convert_document,
    identity_hash,
    verify_converted,
)
from data_center.documents.gateway import (
    ModelStepPending,
    embedding_batches,
    invoke_step,
    reply_vectors,
)
from data_center.documents.locking import database_lock
from data_center.documents.models import (
    DocumentChunkRow,
    DocumentRequestRow,
    DocumentRow,
    DocumentSnapshotRow,
)
from data_center.retrieval.ranking import lexical_score, rank_rrf, select_source_coverage

EMBEDDING_MODEL = "qwen3.7-text-embedding"


def asset_id(source_id: str, record_id: str) -> str:
    return f"{source_id}::{record_id}"


def permitted(scope: AuthorizedScope, source_id: str, record_id: str) -> bool:
    return source_id in scope.source_ids and (
        source_id in scope.source_wide or asset_id(source_id, record_id) in scope.asset_ids
    )


def _chunk(row: DocumentChunkRow) -> DocumentChunk:
    return DocumentChunk(ref=DataEvidenceRef.model_validate(row.evidence), text=row.text,
                         text_sha256=row.text_sha256, metadata=row.metadata_payload)


class DocumentEngine:
    def __init__(self, database: Engine, artifact_root: Path, authorization: AuthorizationPort,
                 source: SourceReadPort, models: ModelGateway, *,
                 conversion_policy: ConversionPolicy | None = None, graph_config=None):
        self.database = database
        self.artifact_root = artifact_root
        self.authorization = authorization
        self.source = source
        self.models = models
        self.conversion_policy = conversion_policy or ConversionPolicy()
        self.graph_config = graph_config

    async def _scope(self, context: CallContext, source_ids: list[str], action: str,
                     expected: AuthorizedScope | None = None) -> AuthorizedScope:
        action = {"publish": "ingest", "model_dispatch": "model"}.get(action, action)
        scope = await self.authorization.authorize(context, source_ids, action)
        if scope.actor_id != context.actor_id or not set(source_ids).issubset(scope.source_ids):
            raise PermissionError("文档来源超出允许范围")
        if expected is not None and (scope.policy_version != expected.policy_version
                                     or scope.fingerprint != expected.fingerprint):
            raise PermissionError("文档操作期间权限版本发生变化")
        return scope

    def _require_documents(self, scope: AuthorizedScope, documents) -> None:
        if any(not permitted(scope, item.source_id, item.record_id) for item in documents):
            raise PermissionError("文档资料超出允许范围")

    def _assert_published(self, scope: AuthorizedScope, snapshot_ids: list[str]) -> None:
        with authorized_connection(self.database, scope) as connection, Session(connection) as session:
            for snapshot_id in snapshot_ids:
                row = session.get(DocumentSnapshotRow, snapshot_id)
                if row is None or row.state != "published":
                    raise PermissionError("资料快照已撤回或不在允许范围")

    def _request(self, context, operation, input_sha256, task_id) -> DocumentRequestRow:
        key = (context.actor_id, context.request_id, operation)
        with Session(self.database) as session, session.begin():
            row = session.get(DocumentRequestRow, key)
            if row is None:
                row = DocumentRequestRow(actor_id=context.actor_id, request_id=context.request_id,
                                         operation=operation, input_sha256=input_sha256, task_id=task_id)
                session.add(row)
                session.flush()
            elif row.input_sha256 != input_sha256 or row.task_id != task_id:
                raise ValueError("相同文档请求编号绑定了其他内容或任务")
            session.expunge(row)
            return row

    async def ingest(self, payload: DocumentIngest, context: CallContext, *,
                     task_id: str, run_id: str) -> DataSnapshot:
        source_ids = sorted({item.source_id for item in payload.documents})
        scope = await self._scope(context, source_ids, "publish")
        self._require_documents(scope, payload.documents)
        keys = [(item.source_id, item.record_id) for item in payload.documents]
        if len(set(keys)) != len(keys):
            raise ValueError("单次发布包含重复文档对象")
        source_revisions = {}
        for item in payload.documents:
            if item.source_id in source_revisions and source_revisions[item.source_id] != item.source_revision:
                raise ValueError("单次发布同一来源包含多个版本")
            source_revisions[item.source_id] = item.source_revision
        input_sha256 = identity_hash({"payload": payload.model_dump(mode="json"),
                                      "conversion": self.conversion_policy.model_dump(mode="json")})
        request_lock = f"dc-document-ingest:{context.actor_id}:{context.request_id}"
        async with database_lock(self.database, request_lock):
            request = self._request(context, "ingest", input_sha256, task_id)
            converted = []
            for item in sorted(payload.documents, key=lambda entry: (entry.source_id, entry.record_id)):
                await self._scope(context, source_ids, "publish", scope)
                raw = await self.source.read_file(item.resource, context)
                converted.append(await asyncio.to_thread(convert_document, item, raw,
                                                          self.artifact_root / "documents",
                                                          self.conversion_policy))
            identity = {"documents": [{"source_id": doc.input.source_id, "record_id": doc.input.record_id,
                                        "revision_id": doc.revision_id, "content_sha256": doc.content_sha256,
                                        "config_sha256": doc.config_sha256} for doc in converted],
                        "embedding": {"model": EMBEDDING_MODEL, "dimensions": 1024},
                        "graph_enabled": payload.graph_enabled}
            snapshot_id = identity_hash(identity)
            if request.snapshot_id is not None and request.snapshot_id != snapshot_id:
                raise ValueError("恢复文档转换与原快照不一致")
            async with database_lock(self.database, f"dc-document-snapshot:{snapshot_id}"):
                with Session(self.database) as session, session.begin():
                    row = session.get(DocumentSnapshotRow, snapshot_id)
                    if row is None:
                        snapshot = DataSnapshot(snapshot_id=snapshot_id, source_ids=source_ids,
                                                source_revisions=source_revisions,
                                                created_at=datetime.now(UTC), state="staged",
                                                record_count=len(converted), schema_sha256=identity_hash(identity["embedding"]),
                                                content_sha256=identity_hash({"documents": identity["documents"]}),
                                                quality={"conversion": "verified", "chunks": sum(len(doc.chunks) for doc in converted)})
                        row = DocumentSnapshotRow(snapshot_id=snapshot_id, state="staged",
                                                  created_at=snapshot.created_at, payload=snapshot.model_dump(mode="json"),
                                                  graph_enabled=payload.graph_enabled)
                        session.add(row)
                        session.flush()
                        self._stage_documents(session, snapshot, converted)
                    elif row.state == "withdrawn":
                        raise ValueError("已撤回快照不能继续发布")
                    else:
                        snapshot = DataSnapshot.model_validate_json(json.dumps(row.payload))
                    request_row = session.get(DocumentRequestRow, (context.actor_id, context.request_id, "ingest"))
                    request_row.snapshot_id = snapshot_id
                if snapshot.state == "published":
                    await self._scope(context, source_ids, "publish", scope)
                    return snapshot
                await self._embed_snapshot(snapshot, scope, context, task_id, run_id)
                await self._scope(context, source_ids, "publish", scope)
                for document in converted:
                    verify_converted(document)
                with Session(self.database) as session, session.begin():
                    rows = list(session.scalars(select(DocumentChunkRow).where(DocumentChunkRow.snapshot_id == snapshot_id)))
                    if not rows or any(item.embedding is None or not item.embedding_call_id for item in rows):
                        raise ValueError("发布索引没有完整真实向量")
                    published = snapshot.model_copy(update={"state": "published"})
                    row = session.get(DocumentSnapshotRow, snapshot_id)
                    row.state, row.payload = "published", published.model_dump(mode="json")
                    request_row = session.get(DocumentRequestRow, (context.actor_id, context.request_id, "ingest"))
                    request_row.result_payload = published.model_dump(mode="json")
                return published

    def _stage_documents(self, session, snapshot, documents):
        for document in documents:
            item = document.input
            identifier = asset_id(item.source_id, item.record_id)
            session.add(DocumentRow(snapshot_id=snapshot.snapshot_id, source_id=item.source_id,
                                    record_id=item.record_id, asset_id=identifier,
                                    source_revision=item.source_revision, source_sha256=document.source_sha256,
                                    content_sha256=document.content_sha256,
                                    converted=document.model_dump(mode="json")))
            for chunk in document.chunks:
                locator = {**chunk.locator, "converter_version": document.converter_version,
                           "conversion_config_sha256": document.config_sha256,
                           "structured_sha256": document.structured_sha256}
                ref = DataEvidenceRef(kind="document", source_id=item.source_id,
                                      snapshot_id=snapshot.snapshot_id, record_id=item.record_id,
                                      source_revision=item.source_revision, source_sha256=document.source_sha256,
                                      content_sha256=document.content_sha256, chunk_id=chunk.chunk_id,
                                      start_line=chunk.start_line, end_line=chunk.end_line,
                                      locator=locator, resource=item.resource)
                session.add(DocumentChunkRow(snapshot_id=snapshot.snapshot_id, chunk_id=chunk.chunk_id,
                                             source_id=item.source_id, record_id=item.record_id,
                                             asset_id=identifier, text=chunk.text, text_sha256=chunk.text_sha256,
                                             evidence=ref.model_dump(mode="json"),
                                             metadata_payload={"asset_id": identifier, "title": item.title,
                                                               "source_nature": item.source_nature,
                                                               "license_id": item.license_id}))

    async def _embed_snapshot(self, snapshot, scope, context, task_id, run_id):
        with Session(self.database) as session:
            rows = list(session.scalars(select(DocumentChunkRow).where(
                DocumentChunkRow.snapshot_id == snapshot.snapshot_id,
                DocumentChunkRow.embedding.is_(None)).order_by(DocumentChunkRow.chunk_id)))
            for row in rows:
                session.expunge(row)
        pending = {}
        for row in rows:
            item = pending.setdefault(row.text_sha256, {"text": row.text, "chunk_ids": []})
            if item["text"] != row.text:
                raise ValueError("embedding文字摘要冲突")
            item["chunk_ids"].append(row.chunk_id)
        entries = [pending[key] for key in sorted(pending)]
        by_text = {entry["text"]: entry for entry in entries}
        for texts in embedding_batches([entry["text"] for entry in entries]):
            selected = [by_text[text] for text in texts]
            await self._scope(context, snapshot.source_ids, "model_dispatch", scope)
            step_key = "document-embedding:" + identity_hash({"snapshot": snapshot.snapshot_id,
                                                           "texts": texts, "policy": scope.policy_version})
            request = ModelInvocation(task_id=task_id, run_id=run_id, step_key=step_key,
                                      kind="embedding", model=EMBEDDING_MODEL, texts=texts,
                                      source_ids=snapshot.source_ids)
            reply = await invoke_step(self.models, request, context)
            vectors = reply_vectors(reply, len(selected))
            await self._scope(context, snapshot.source_ids, "publish", scope)
            with Session(self.database) as session, session.begin():
                for entry, vector in zip(selected, vectors, strict=True):
                    for chunk_id in entry["chunk_ids"]:
                        row = session.get(DocumentChunkRow, (snapshot.snapshot_id, chunk_id))
                        row.embedding, row.embedding_call_id = vector, reply.request_id

    async def _visible(self, query, context, snapshot_ids=None):
        if query.intent != "documents":
            raise ValueError("DocumentEngine只接受文档查询")
        scope = await self._scope(context, query.source_ids, "read")
        with authorized_connection(self.database, scope, query.snapshot_id) as connection, Session(connection) as session:
            statement = select(DocumentSnapshotRow).where(DocumentSnapshotRow.state == "published")
            if query.snapshot_id:
                statement = statement.where(DocumentSnapshotRow.snapshot_id == query.snapshot_id)
            elif snapshot_ids is not None:
                statement = statement.where(DocumentSnapshotRow.snapshot_id.in_(snapshot_ids))
            candidates = list(session.scalars(statement.order_by(DocumentSnapshotRow.created_at.desc(),
                                                                  DocumentSnapshotRow.snapshot_id)))
            chosen = {}
            source_snapshots = {}
            for row in candidates:
                snapshot = DataSnapshot.model_validate_json(json.dumps(row.payload))
                for source_id in query.source_ids:
                    if source_id in snapshot.source_ids and source_id not in source_snapshots:
                        chosen[row.snapshot_id] = snapshot
                        source_snapshots[source_id] = row.snapshot_id
            rows = []
            for source_id, snapshot_id in source_snapshots.items():
                statement = select(DocumentChunkRow).where(DocumentChunkRow.snapshot_id == snapshot_id,
                                                            DocumentChunkRow.source_id == source_id)
                if source_id not in scope.source_wide:
                    statement = statement.where(DocumentChunkRow.asset_id.in_(scope.asset_ids))
                if query.record_id:
                    statement = statement.where(DocumentChunkRow.record_id == query.record_id)
                rows.extend(session.scalars(statement.order_by(DocumentChunkRow.chunk_id)))
            for row in rows:
                if not permitted(scope, row.source_id, row.record_id):
                    raise PermissionError("文档候选超出允许资料集合")
                session.expunge(row)
        await self._scope(context, query.source_ids, "read", scope)
        visible_snapshots = []
        for snapshot in chosen.values():
            visible = [row for row in rows if row.snapshot_id == snapshot.snapshot_id]
            if not visible:
                continue
            visible_sources = sorted({row.source_id for row in visible})
            visible_snapshots.append(snapshot.model_copy(update={
                "source_ids": visible_sources,
                "source_revisions": {source_id: snapshot.source_revisions[source_id] for source_id in visible_sources},
                "record_count": len({row.asset_id for row in visible}),
                "quality": {"authorized_scope": scope.fingerprint, "chunks": len(visible)},
            }))
        return scope, visible_snapshots, rows

    async def search(self, query: DataQuery, context: CallContext, *,
                     task_id: str, run_id: str) -> DocumentSearchResult:
        if query.intent != "documents":
            raise ValueError("DocumentEngine只接受文档查询")
        if query.filters:
            raise ValueError("文档查询不支持未登记的filters")
        query_payload = query.model_dump(mode="json")
        input_sha256 = identity_hash({"query": query_payload})
        async with database_lock(self.database, f"dc-document-search:{context.actor_id}:{context.request_id}"):
            scope = await self._scope(context, query.source_ids, "read")
            key = (context.actor_id, context.request_id, "search")
            with Session(self.database) as session:
                request = session.get(DocumentRequestRow, key)
                if request is not None:
                    session.expunge(request)
            envelope = None
            cached = None
            if request is not None:
                if request.task_id != task_id:
                    raise ValueError("相同文档请求编号绑定了其他内容或任务")
                stored = request.result_payload
                if stored is not None and stored.get("binding_version") == 1:
                    if request.input_sha256 != input_sha256 or stored["query"] != query_payload:
                        raise ValueError("相同文档请求编号绑定了其他内容或任务")
                    envelope = stored
                    cached = DocumentSearchResult.model_validate_json(json.dumps(stored["result"]))
                else:
                    legacy_sha256 = identity_hash({"query": query_payload,
                                                    "scope": scope.fingerprint, "policy": scope.policy_version})
                    if request.input_sha256 != legacy_sha256:
                        raise ValueError("历史查询请求的内容与原授权版本不能确认")
                    if stored is not None:
                        cached = DocumentSearchResult.model_validate_json(json.dumps(stored))
            bound_ids = ([item.snapshot_id for item in cached.snapshots] if cached is not None else None)
            scope, snapshots, rows = await self._visible(query, context, bound_ids)
            if cached is not None and set(bound_ids) != {item.snapshot_id for item in snapshots}:
                raise PermissionError("原查询快照已撤回或不在允许范围")
            if envelope is None:
                envelope = {"binding_version": 1, "query": query_payload,
                            "scope": scope.fingerprint, "policy": scope.policy_version,
                            "allowed_chunk_ids": sorted(f"{row.snapshot_id}:{row.chunk_id}" for row in rows),
                            "result": (cached or DocumentSearchResult(status="waiting", snapshots=snapshots)).model_dump(mode="json")}
                with Session(self.database) as session, session.begin():
                    row = session.get(DocumentRequestRow, key)
                    if row is None:
                        row = DocumentRequestRow(actor_id=context.actor_id, request_id=context.request_id,
                                                 operation="search", input_sha256=input_sha256, task_id=task_id)
                        session.add(row)
                    row.input_sha256, row.result_payload = input_sha256, envelope
            if cached is not None:
                original_ids = set(envelope["allowed_chunk_ids"])
                current_snapshots = []
                for snapshot in cached.snapshots:
                    visible = [row for row in rows if row.snapshot_id == snapshot.snapshot_id
                               and f"{row.snapshot_id}:{row.chunk_id}" in original_ids]
                    sources = sorted({row.source_id for row in visible})
                    current_snapshots.append(snapshot.model_copy(update={
                        "source_ids": sources,
                        "source_revisions": {source_id: snapshot.source_revisions[source_id] for source_id in sources},
                        "record_count": len({row.asset_id for row in visible}),
                        "quality": {**snapshot.quality, "chunks": len(visible)},
                    }))
                snapshots = current_snapshots
            if cached is not None and cached.status == "completed":
                for hit in cached.hits:
                    await self.read_evidence(hit.chunk.ref, context)
                if query.retrieval_mode == "graph":
                    allowed = {f"{row.snapshot_id}:{row.chunk_id}" for row in rows}
                    if not set(envelope["allowed_chunk_ids"]).issubset(allowed):
                        raise PermissionError("原图查询资料超出当前允许范围")
                    from data_center.graph.engine import GraphEngine

                    await GraphEngine(self, self.graph_config).validate_cached(cached, rows, context)
                await self._scope(context, query.source_ids, "read", scope)
                if query.required_source_ids:
                    coverage = self._coverage_metadata(query, [hit.chunk.ref.source_id for hit in cached.hits])
                    if coverage["missing_source_ids"]:
                        blocked = cached.model_copy(update={
                            "status": "DATA_BLOCKED", "snapshots": snapshots, "hits": [],
                            "metadata": {**cached.metadata, **coverage, "reason": "必需文档来源没有可回读引用"},
                        })
                        with Session(self.database) as session, session.begin():
                            session.get(DocumentRequestRow, key).result_payload = {
                                **envelope, "result": blocked.model_dump(mode="json"),
                            }
                        return blocked
                return cached.model_copy(update={"snapshots": snapshots})
            if cached is not None and cached.status in {"unknown", "failed", "cancelled", "DATA_BLOCKED"}:
                return cached.model_copy(update={"snapshots": snapshots})
            if cached is not None and (envelope["scope"] != scope.fingerprint
                                       or envelope["policy"] != scope.policy_version):
                raise PermissionError("恢复查询期间授权版本发生变化")
            if query.required_source_ids:
                coverage = self._coverage_metadata(query, [row.source_id for row in rows])
                if coverage["missing_source_ids"]:
                    result = DocumentSearchResult(status="DATA_BLOCKED", snapshots=snapshots,
                        metadata={**coverage, "reason": "必需文档来源没有可回读片段"})
                    with Session(self.database) as session, session.begin():
                        session.get(DocumentRequestRow, key).result_payload = {
                            **envelope, "result": result.model_dump(mode="json"),
                        }
                    return result
            if not rows:
                result = DocumentSearchResult(status="DATA_BLOCKED", metadata={"reason": "授权范围内没有已发布文档"})
                envelope = {**envelope, "result": result.model_dump(mode="json")}
                with Session(self.database) as session, session.begin():
                    session.get(DocumentRequestRow, key).result_payload = envelope
                return result
            if query.retrieval_mode == "graph":
                from data_center.graph.engine import GraphEngine

                result = await GraphEngine(self, self.graph_config).search(query, scope, snapshots, rows,
                                                                           context, task_id, run_id)
                if query.required_source_ids and result.status == "completed":
                    coverage = self._coverage_metadata(query, [hit.chunk.ref.source_id for hit in result.hits])
                    supplement = None
                    if coverage["missing_source_ids"]:
                        supplement_query = query.model_copy(update={
                            "required_source_ids": coverage["missing_source_ids"],
                        })
                        supplement = await self._hybrid(supplement_query, scope, snapshots, rows,
                                                        context, task_id, run_id)
                    result = self._merge_source_coverage(query, result, supplement)
            else:
                result = await self._hybrid(query, scope, snapshots, rows, context, task_id, run_id)
            await self._scope(context, query.source_ids, "read", scope)
            with Session(self.database) as session, session.begin():
                row = session.get(DocumentRequestRow, (context.actor_id, context.request_id, "search"))
                row.result_payload = {**envelope, "result": result.model_copy(update={"snapshots": snapshots}).model_dump(mode="json")}
            return result

    async def _hybrid(self, query, scope, snapshots, rows, context, task_id, run_id):
        await self._scope(context, query.source_ids, "model_dispatch", scope)
        self._assert_published(scope, [snapshot.snapshot_id for snapshot in snapshots])
        step_key = "document-query:" + identity_hash({"question": query.question,
                                                      "snapshots": [item.snapshot_id for item in snapshots],
                                                      "scope": scope.fingerprint, "policy": scope.policy_version})
        request = ModelInvocation(task_id=task_id, run_id=run_id, step_key=step_key,
                                  kind="embedding", model=EMBEDDING_MODEL,
                                  texts=[query.question], source_ids=query.source_ids)
        try:
            reply = await invoke_step(self.models, request, context)
        except ModelStepPending as pending:
            status = pending.reply.status if pending.reply.status in {"unknown", "failed", "cancelled"} else "waiting"
            return DocumentSearchResult(status=status, snapshots=snapshots,
                                        call_ids=[pending.reply.request_id])
        vector = reply_vectors(reply, 1)[0]
        await self._scope(context, query.source_ids, "read", scope)
        ranked, missing = self._rank_hybrid(query, scope, rows, vector)
        if missing:
            coverage = self._coverage_metadata(query, [row.source_id for row in rows])
            return DocumentSearchResult(status="DATA_BLOCKED", snapshots=snapshots,
                call_ids=[reply.request_id], metadata={**coverage, "missing_source_ids": missing,
                                                       "reason": "必需文档来源没有可回读候选"})
        indexed = {f"{row.snapshot_id}:{row.chunk_id}": row for row in rows}
        hits = []
        for rank, entry in enumerate(ranked, 1):
            chunk = await self.read_evidence(DataEvidenceRef.model_validate(indexed[entry["key"]].evidence), context)
            hits.append(DocumentSearchHit(rank=rank, chunk=chunk,
                                          lexical_score=entry["lexical_score"],
                                          vector_distance=entry["vector_distance"], rrf_score=entry["rrf_score"]))
        metadata = {"candidate_count": len(rows), "retrieval_mode": "hybrid", "rrf_k": 60}
        if query.required_source_ids:
            metadata.update(self._coverage_metadata(query, [hit.chunk.ref.source_id for hit in hits]))
            metadata["coverage_supplements"] = [
                {"snapshot_id": indexed[item["key"]].snapshot_id, "chunk_id": indexed[item["key"]].chunk_id,
                 "source_id": indexed[item["key"]].source_id,
                 "retrieval_mode": "hybrid", "retrieval_scope": "source"}
                for item in ranked if item.get("retrieval_scope") == "source"
            ]
        return DocumentSearchResult(status="completed", snapshots=snapshots, hits=hits,
                                    call_ids=[reply.request_id], metadata=metadata)

    @staticmethod
    def _coverage_metadata(query, source_ids):
        required = set(query.required_source_ids)
        covered = required & set(source_ids)
        return {"required_source_ids": list(query.required_source_ids),
                "covered_source_ids": sorted(covered), "missing_source_ids": sorted(required - covered)}

    def _merge_source_coverage(self, query, primary, supplement=None):
        if not query.required_source_ids:
            return primary
        call_ids = list(dict.fromkeys([*primary.call_ids, *(supplement.call_ids if supplement else [])]))
        if supplement is not None and supplement.status != "completed":
            coverage = self._coverage_metadata(query, [hit.chunk.ref.source_id for hit in primary.hits])
            return primary.model_copy(update={"status": supplement.status, "call_ids": call_ids,
                "metadata": {**primary.metadata, **coverage,
                             "coverage_supplement_status": supplement.status,
                             "coverage_supplement_metadata": supplement.metadata}})
        primary_keys = {f"{hit.chunk.ref.snapshot_id}:{hit.chunk.ref.chunk_id}" for hit in primary.hits}
        hits = {}
        for hit in [*primary.hits, *(supplement.hits if supplement else [])]:
            ref = hit.chunk.ref
            key = f"{ref.snapshot_id}:{ref.chunk_id}"
            hits.setdefault(key, hit)
        ranked = [{"key": key} for key in hits]
        source_by_key = {key: hit.chunk.ref.source_id for key, hit in hits.items()}
        selected, missing = select_source_coverage(ranked, source_by_key, query.required_source_ids,
                                                   limit=query.limit)
        coverage = self._coverage_metadata(query, source_by_key.values())
        if missing:
            return primary.model_copy(update={"status": "DATA_BLOCKED", "hits": [], "call_ids": call_ids,
                "metadata": {**primary.metadata, **coverage, "reason": "必需文档来源没有可回读引用"}})
        chosen = [hits[item["key"]].model_copy(update={"rank": rank}) for rank, item in enumerate(selected, 1)]
        supplements = [{"snapshot_id": hit.chunk.ref.snapshot_id, "chunk_id": hit.chunk.ref.chunk_id,
                        "source_id": hit.chunk.ref.source_id, "retrieval_mode": "hybrid"}
                       for item, hit in zip(selected, chosen, strict=True) if item["key"] not in primary_keys]
        coverage = self._coverage_metadata(query, [hit.chunk.ref.source_id for hit in chosen])
        return primary.model_copy(update={"hits": chosen, "call_ids": call_ids,
            "metadata": {**primary.metadata, **coverage, "coverage_supplements": supplements}})

    def _vector_ranking(self, scope, rows, vector, pool_size):
        if any(row.embedding is None for row in rows):
            raise ValueError("已发布文档缺少真实embedding")
        with authorized_connection(self.database, scope) as connection:
            distance = DocumentChunkRow.embedding.cosine_distance(vector).label("distance")
            statement = select(DocumentChunkRow.snapshot_id, DocumentChunkRow.chunk_id, distance).where(
                tuple_(DocumentChunkRow.snapshot_id, DocumentChunkRow.chunk_id).in_(
                    [(row.snapshot_id, row.chunk_id) for row in rows]),
            ).order_by(distance, DocumentChunkRow.snapshot_id, DocumentChunkRow.chunk_id).limit(pool_size)
            return [(f"{snapshot_id}:{chunk_id}", score)
                    for snapshot_id, chunk_id, score in connection.execute(statement)]

    def _rank_hybrid(self, query, scope, rows, vector):
        pool_size = max(query.limit * 10, 100)
        indexed = {f"{row.snapshot_id}:{row.chunk_id}": row for row in rows}
        lexical_scores = {key: lexical_score(query.question, row.text) for key, row in indexed.items()}
        lexical = sorted((entry for entry in lexical_scores.items() if entry[1] > 0),
                         key=lambda entry: (-entry[1], entry[0]))[:pool_size]
        distances = self._vector_ranking(scope, rows, vector, pool_size)
        ranked = rank_rrf(lexical, distances, limit=query.limit)
        if not query.required_source_ids:
            return ranked, []
        source_by_key = {key: row.source_id for key, row in indexed.items()}
        covered = {source_by_key[item["key"]] for item in ranked}
        for source_id in sorted(set(query.required_source_ids) - covered):
            source_rows = [row for row in rows if row.source_id == source_id]
            if not source_rows:
                continue
            source_keys = {f"{row.snapshot_id}:{row.chunk_id}" for row in source_rows}
            source_lexical = sorted([(key, lexical_scores[key]) for key in source_keys if lexical_scores[key] > 0],
                                    key=lambda item: (-item[1], item[0]))[:pool_size]
            source_distances = self._vector_ranking(scope, source_rows, vector, pool_size)
            source_ranked = rank_rrf(source_lexical, source_distances, limit=1)
            ranked.extend({**item, "retrieval_scope": "source"} for item in source_ranked)
        return select_source_coverage(ranked, source_by_key, query.required_source_ids, limit=query.limit)

    async def _read_verified_document(self, ref: DataEvidenceRef, context: CallContext
                                      ) -> tuple[DocumentChunk, ConvertedDocument, AuthorizedScope]:
        if ref.kind != "document":
            raise ValueError("DocumentEngine只接受文档引用")
        scope = await self._scope(context, [ref.source_id], "read")
        if not permitted(scope, ref.source_id, ref.record_id):
            raise PermissionError("文档引用超出允许资料集合")
        with authorized_connection(self.database, scope, ref.snapshot_id) as connection, Session(connection) as session:
            snapshot = session.get(DocumentSnapshotRow, ref.snapshot_id)
            row = session.get(DocumentChunkRow, (ref.snapshot_id, ref.chunk_id))
            document = session.get(DocumentRow, (ref.snapshot_id, ref.source_id, ref.record_id))
            if snapshot is None or snapshot.state != "published" or row is None or document is None:
                raise PermissionError("文档引用不存在或已撤回")
            identity = DataSnapshot.model_validate_json(json.dumps(snapshot.payload))
            if ref.source_id not in identity.source_ids or identity.source_revisions[ref.source_id] != ref.source_revision:
                raise ValueError("文档引用与历史快照来源版本不一致")
            chunk = _chunk(row)
            converted = ConvertedDocument.model_validate(document.converted)
        if chunk.ref != ref:
            raise ValueError("文档引用身份或位置不一致")
        verify_converted(converted)
        if (converted.input.source_id != ref.source_id or converted.input.record_id != ref.record_id
                or converted.input.source_revision != ref.source_revision
                or converted.input.resource != ref.resource
                or converted.source_sha256 != ref.source_sha256
                or converted.content_sha256 != ref.content_sha256):
            raise ValueError("文档引用与不可变原件身份不一致")
        original_chunk = next((item for item in converted.chunks if item.chunk_id == ref.chunk_id), None)
        if (original_chunk is None or original_chunk.start_line != ref.start_line
                or original_chunk.end_line != ref.end_line
                or original_chunk.text_sha256 != chunk.text_sha256
                or any(ref.locator.get(key) != value for key, value in original_chunk.locator.items())):
            raise ValueError("文档引用与结构化原件位置不一致")
        lines = converted.text.splitlines(keepends=True)
        content = "".join(lines[ref.start_line - 1:ref.end_line])
        if content != chunk.text or hashlib.sha256(content.encode()).hexdigest() != chunk.text_sha256:
            raise ValueError("文档引用回读内容不一致")
        await self._scope(context, [ref.source_id], "read", scope)
        return chunk, converted, scope

    async def read_evidence(self, ref: DataEvidenceRef, context: CallContext) -> DocumentChunk:
        chunk, _, _ = await self._read_verified_document(ref, context)
        return chunk

    async def read_document_statistics(self, ref: DataEvidenceRef, context: CallContext) -> dict:
        _, converted, scope = await self._read_verified_document(ref, context)
        normalized = " ".join(converted.text.split())
        statistics = {"reference": ref.model_dump(mode="json"), "source_id": ref.source_id,
                      "snapshot_id": ref.snapshot_id, "record_id": ref.record_id,
                      "source_revision": ref.source_revision, "source_sha256": converted.source_sha256,
                      "content_sha256": converted.content_sha256,
                      "conversion_config_sha256": converted.config_sha256,
                      "structured_sha256": converted.structured_sha256,
                      "converter_version": converted.converter_version,
                      "resource": converted.input.resource.model_dump(mode="json"),
                      "title": converted.input.title, "document_scope": "complete_saved_conversion",
                      "normalization": "whitespace_collapsed",
                      "normalization_rule": "使用Python str.split()识别Unicode空白，再用单个U+0020 SPACE连接；保留标点、数字及其余字符",
                      "normalized_character_count_unit": "Unicode code points",
                      "normalized_character_count": len(normalized), "normalized_text": normalized,
                      "normalized_text_sha256": hashlib.sha256(normalized.encode()).hexdigest()}
        if converted.original_path.suffix.casefold() == ".pptx":
            original = converted.original_path.read_bytes()
            if hashlib.sha256(original).hexdigest() != converted.source_sha256:
                raise ValueError("统计读取的PPTX原件摘要不一致")
            statistics["page_count"] = len(Presentation(BytesIO(original)).slides)
            statistics["page_count_method"] = "python-pptx.Presentation.slides"
        if converted.original_path.suffix.casefold() == ".xlsx":
            original = converted.original_path.read_bytes()
            if hashlib.sha256(original).hexdigest() != converted.source_sha256:
                raise ValueError("统计读取的XLSX原件摘要不一致")
            workbook = load_workbook(BytesIO(original), read_only=False, data_only=False)
            try:
                statistics["worksheet_reader"] = "openpyxl.load_workbook"
                statistics["workbook_scope"] = "complete_native_workbook"
                statistics["workbook_data_only"] = False
                statistics["worksheet_count"] = len(workbook.worksheets)
                statistics["worksheets"] = [
                    {"worksheet": sheet.title, "cell_range": sheet.calculate_dimension(),
                     "min_row": sheet.min_row, "max_row": sheet.max_row,
                     "min_column": sheet.min_column, "max_column": sheet.max_column,
                     "merged_cell_ranges": sorted(str(value) for value in sheet.merged_cells.ranges),
                     "rows": [{"row_number": cells[0].row,
                               "cell_range": f"{cells[0].coordinate}:{cells[-1].coordinate}",
                               "cells": [{"coordinate": cell.coordinate, "row": cell.row,
                                          "column": cell.column, "data_type": cell.data_type,
                                          "value": to_jsonable_python(cell.value)} for cell in cells]}
                              for cells in sheet.iter_rows()]}
                    for sheet in workbook.worksheets]
            finally:
                workbook.close()
        await self._scope(context, [ref.source_id], "read", scope)
        return statistics

    async def build_graph(self, query: DataQuery, context: CallContext, *, task_id: str, run_id: str) -> dict:
        from data_center.graph.engine import GraphEngine

        scope, snapshots, rows = await self._visible(query, context)
        if query.required_source_ids:
            coverage = self._coverage_metadata(query, [row.source_id for row in rows])
            if coverage["missing_source_ids"]:
                return {"state": "DATA_BLOCKED", **coverage, "reason": "必需文档来源没有可回读片段"}
        return await GraphEngine(self, self.graph_config).build(query, scope, snapshots, rows,
                                                               context, task_id, run_id)

    async def graph_status(self, query: DataQuery, context: CallContext) -> dict:
        from data_center.graph.engine import GraphEngine

        scope, snapshots, rows = await self._visible(query, context)
        graph = GraphEngine(self, self.graph_config)
        await graph.resolve_identity(query, scope, snapshots, rows, context)
        return graph.status(scope, snapshots, rows)
